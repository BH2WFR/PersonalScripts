"""Verify selective archive loads and numeric Excel/text range round trips.

Requirements: viewer dependencies, including openpyxl. Usage: unittest discovery.
Generated workbooks stay in the Git-ignored automated fixture directory.
"""

from dataclasses import replace
import importlib
from io import BytesIO
from unittest import TestCase
from unittest.mock import patch

import h5py
import numpy as np
from openpyxl import Workbook, load_workbook
from scipy.io import savemat

from fixture_store import fixture_directory
from test_data_model import model

catalog = importlib.import_module("personal_matrix_viewer.import_catalog")
excel = importlib.import_module("personal_matrix_viewer.excel_io")
exports = importlib.import_module("personal_matrix_viewer.exporting")
tables = importlib.import_module("personal_matrix_viewer.table_data")
workspace = importlib.import_module("personal_matrix_viewer.workspace")


class TableImportTests(TestCase):
    """Use actual files to check numeric fidelity and transactional selection."""

    def setUp(self) -> None:
        """Place generated files in the ignored scratch tree, never manual samples."""
        self.folder = fixture_directory(f"unit-files/table-import/{self._testMethodName}")

    def test_text_range_excludes_preamble_labels_and_footer(self) -> None:
        """Original row/column coordinates select data before type validation."""
        for suffix, delimiter in (("csv", ","), ("txt", "\t")):
            path = self.folder / f"ranged.{suffix}"
            path.write_text("Experiment metadata\n" + "\n".join(delimiter.join(row) for row in (
                ("label", "x", "y", "notes"), ("a", "1", "10", "keep out"),
                ("b", "2", "20", "keep out"), ("footer", "bad", "bad", "bad"))) + "\n", encoding="utf-8")
            inspected = catalog.inspect_source(path)
            self.assertEqual(inspected.members[0].shape, (5, 4))
            region = tables.TableRange(2, 2, 4, 3)
            doc = model.load_document(path, region=region, header=tables.HeaderMode.FIRST)
            np.testing.assert_array_equal(doc.array, [[1, 10], [2, 20]])
            self.assertEqual(doc.csv_headers, ("x", "y"))
            self.assertTrue(doc.import_provenance)
            with self.assertRaisesRegex(ValueError, "row 5, column 2"):
                model.load_document(path, region=replace(region, row_end=5), header=tables.HeaderMode.FIRST)
            with self.assertRaises(ValueError):
                model.load_document(path, region=tables.TableRange(6))
            with self.assertRaises(ValueError):
                model.load_document(path, delimiter=",;")

    def test_text_integer_header_none_and_blank_records(self) -> None:
        """An explicit numeric header and whitespace records do not corrupt data."""
        path = self.folder / "integers.csv"
        path.write_text("2026\n \n9007199254740993\n18446744073709551615\n", encoding="utf-8")
        doc = model.load_document(path, header=tables.HeaderMode.FIRST)
        np.testing.assert_array_equal(doc.array, np.array([[2**53 + 1], [2**64 - 1]], dtype=np.uint64))
        complete = model.load_document(path, header=tables.HeaderMode.NONE)
        self.assertEqual(complete.array.shape, (3, 1))

    def test_npz_catalog_disables_unsupported_and_loads_only_selected(self) -> None:
        """Inspect headers without unpickling or allocating unselected matrices."""
        path = self.folder / "members.npz"
        source = np.arange(20).reshape(4, 5)
        np.savez(path, selected=source, other=np.arange(7), objects=np.array([{}], dtype=object), cube=np.zeros((2, 2, 2)))
        inspected = catalog.inspect_source(path)
        self.assertEqual([member.shape for member in inspected.members], [(4, 5), (7,), (1,), (2, 2, 2)])
        self.assertEqual([bool(member.error) for member in inspected.members], [False, False, True, True])
        with patch.object(workspace, "load_document", wraps=model.load_document) as loader:
            loaded = workspace.load_file(path, choices=(catalog.ImportChoice("selected"),))
        self.assertEqual(loader.call_count, 1)
        self.assertEqual(len(loaded.documents), 1)
        np.testing.assert_array_equal(loaded.documents[0].array, source)
        with self.assertRaisesRegex(ValueError, "unavailable"):
            catalog.inspect_source(path, "absent")

    def test_mat_catalog_dimensions_and_selection(self) -> None:
        """Both MAT generations use source-visible shape and selective loading."""
        source = np.arange(20).reshape(4, 5)
        legacy = self.folder / "legacy.mat"
        savemat(legacy, {"numeric": source, "other": np.arange(6), "words": "not numeric"})
        modern = self.folder / "modern.mat"
        with h5py.File(modern, "w") as archive:
            item = archive.create_dataset("numeric", data=source.T.astype(np.float64))
            item.attrs["MATLAB_class"] = np.bytes_("double")
            words = archive.create_dataset("words", data=np.array([[65]], dtype=np.uint16))
            words.attrs["MATLAB_class"] = np.bytes_("char")
        for path in (legacy, modern):
            inspected = catalog.inspect_source(path)
            numeric = next(item for item in inspected.members if item.key == "numeric")
            self.assertEqual(numeric.shape, (4, 5))
            self.assertFalse(numeric.error)
            self.assertTrue(next(item for item in inspected.members if item.key == "words").error)
            loaded = workspace.load_file(path, choices=(catalog.ImportChoice("numeric"),))
            np.testing.assert_array_equal(loaded.documents[0].array, source)

    def test_excel_per_sheet_ranges_and_unselected_bad_cells(self) -> None:
        """Independent worksheet rectangles exclude headings and invalid notes."""
        path = self.folder / "ranges.xlsx"
        book = Workbook()
        first = book.active
        assert first is not None
        first.title = "A"
        first.append(["Experiment"])
        first.append(["labels", "x", "y"])
        first.append(["a", 1, 10])
        first.append(["b", 2, 20])
        second = book.create_sheet("B")
        second.append(["ignore", 8, 7])
        second.append(["ignore", 6, 5])
        book.create_sheet("Invalid").append(["=1+2"])
        book.save(path)
        book.close()
        inspected = catalog.inspect_source(path)
        self.assertEqual([member.shape for member in inspected.members], [(4, 3), (2, 3), (1, 1)])
        choices = (catalog.ImportChoice("A", tables.TableRange(2, 2), tables.HeaderMode.FIRST),
                   catalog.ImportChoice("B", tables.TableRange(1, 2), tables.HeaderMode.NONE))
        loaded = workspace.load_file(path, choices=choices)
        self.assertEqual([doc.key for doc in loaded.documents], ["A", "B"])
        np.testing.assert_array_equal(loaded.documents[0].array, [[1, 10], [2, 20]])
        np.testing.assert_array_equal(loaded.documents[1].array, [[8, 7], [6, 5]])
        self.assertEqual(loaded.documents[0].csv_headers, ("x", "y"))
        with self.assertRaisesRegex(ValueError, "formula has no cached result"):
            model.load_document(path, "Invalid")
        with self.assertRaisesRegex(ValueError, "exceeds"):
            model.load_document(path, "A", region=tables.TableRange(1, 4))

    def test_excel_blank_vector_cells_retain_positions(self) -> None:
        """An empty cell between two values is a gap, not a deleted sample."""
        path = self.folder / "gaps.xlsx"
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        sheet.append([1])
        sheet.append([None])
        sheet.append([3])
        book.save(path)
        book.close()
        doc = model.load_document(path)
        np.testing.assert_array_equal(doc.array, [[1], [np.nan], [3]])
        self.assertEqual(model.default_selection(doc).mode, model.ViewMode.SIGNAL)

    def test_empty_sheet_is_rejected_but_explicit_nan_is_supported(self) -> None:
        """An unused worksheet is not silently treated as a one-cell signal."""
        path = self.folder / "empty.xlsx"
        book = Workbook()
        sheet = book.active
        assert sheet is not None
        book.save(path)
        with self.assertRaisesRegex(ValueError, "no numeric records"):
            model.load_document(path)
        sheet.append(["nan"])
        book.save(path)
        book.close()
        np.testing.assert_array_equal(model.load_document(path).array, [[np.nan]])

    def test_excel_real_complex_and_integer_export_round_trip(self) -> None:
        """Separate complex sheets preserve both components and large integers."""
        source = np.array([[1 + 2j, 3 - 4j], [5 + 6j, 7 - 8j]])
        integers = np.array([2**53 + 1, 2**64 - 1], dtype=np.uint64)
        path = self.folder / "export.xlsx"
        path.write_bytes(excel.serialize_excel((("signal" * 8, source), ("large", integers),
                                               ("gaps", np.array([np.nan, np.inf, -np.inf, 0.25])))))
        names = tuple(member.key for member in catalog.inspect_source(path).members)
        self.assertTrue(names[0].endswith("_real"))
        self.assertTrue(names[1].endswith("_imag"))
        real, imag = (model.load_document(path, name).array for name in names[:2])
        np.testing.assert_array_equal(real + 1j * imag, source)
        np.testing.assert_array_equal(model.load_document(path, "large").array[:, 0], integers)
        np.testing.assert_array_equal(model.load_document(path, "gaps").array[:, 0], [np.nan, np.inf, -np.inf, .25])
        self.assertEqual(exports.output_paths(path, source, exports.ExportFormat.XLSX), (path,))
        book = load_workbook(BytesIO(exports.serialize_array(source, exports.ExportFormat.XLSX)))
        try:
            self.assertEqual(book.sheetnames, ["matrix_real", "matrix_imag"])
        finally:
            book.close()

    def test_excel_names_and_limits(self) -> None:
        """Sanitize worksheet names and reject grids exceeding Excel's limits."""
        names = excel.excel_sheet_names(("a/b", "A:B", "'", "x" * 40, "x" * 40))
        self.assertEqual(len({name.casefold() for name in names}), len(names))
        self.assertTrue(all(0 < len(name) <= 31 for name in names))
        for values in (np.zeros((1, 16385)), np.zeros((1, 1, 1)), np.array(["bad"])):
            with self.assertRaises(ValueError):
                excel.serialize_excel((("bad", values),))

    def test_magnitude_db_floor_crop_and_source_preservation(self) -> None:
        """Magnitude dB is a relative display channel, not destructive conversion."""
        source = np.array([1 + 0j, .1j, .001 + 0j, 0j, complex(np.nan, 0)])
        doc = model.Document(self.folder / "complex.npy", source)
        selection = replace(model.default_selection(doc), component=model.Component.MAGNITUDE_DB, db_floor=-40)
        frame = model.prepare_frame(doc, selection, model.Limits(), 2)
        np.testing.assert_allclose(frame.scalar, [0, -20, -40, -40, np.nan])
        cropped = model.prepare_frame(doc, selection, model.Limits(), 2, model.Crop(1, 3))
        np.testing.assert_allclose(cropped.scalar, [0, -40, -40])
        np.testing.assert_array_equal(doc.array, source)
