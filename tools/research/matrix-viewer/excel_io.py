"""Read worksheet rectangles and export numeric XLSX workbooks.

Requirements: numpy and openpyxl; xlrd for legacy XLS input.
Usage: inspect_excel/read_excel for import;
serialize_excel for new workbooks. XLSM reads data only; no macros are executed.
Formula cells require cached results saved by a spreadsheet application.
"""

from collections.abc import Iterator, Sequence
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re
from typing import TYPE_CHECKING, cast

import numpy as np
from numpy.typing import NDArray

from .table_data import HeaderMode, TableData, TableRange, numeric_table

if TYPE_CHECKING:
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    from xlrd.book import Book
    from xlrd.sheet import Sheet

EXCEL_EXTENSIONS = frozenset({".xlsx", ".xlsm", ".xls"})
EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_COLUMNS = 16_384
PREVIEW_ROWS = 8
PREVIEW_COLUMNS = 8


@dataclass(frozen=True)
class ExcelSheet:
    """Worksheet name, declared used rectangle and a small unconverted preview."""

    name: str
    shape: tuple[int, int]
    preview: tuple[tuple[str, ...], ...]


def _dimensions(sheet: "ReadOnlyWorksheet") -> tuple[int, int]:
    if sheet.max_row is None or sheet.max_column is None:
        sheet.calculate_dimension(force=True)
    return int(sheet.max_row or 1), int(sheet.max_column or 1)


def _open_xls(path: Path) -> "Book":
    """Open a legacy workbook lazily, with an actionable missing-library error."""
    try:
        import xlrd
    except ImportError as exc:
        raise ImportError("Reading .xls requires xlrd. Install with: conda run -n base python -m pip install xlrd") from exc
    return xlrd.open_workbook(str(path), on_demand=True)


def _inspect_xls(path: Path) -> tuple[ExcelSheet, ...]:
    """Inspect one sheet at a time, releasing its parsed cells after previewing."""
    workbook = _open_xls(path)
    try:
        sheets: list[ExcelSheet] = []
        for index, name in enumerate(workbook.sheet_names()):
            sheet = workbook.sheet_by_index(index)
            preview = tuple(tuple(str(sheet.cell_value(row, column))
                                  for column in range(min(sheet.ncols, PREVIEW_COLUMNS)))
                            for row in range(min(sheet.nrows, PREVIEW_ROWS)))
            sheets.append(ExcelSheet(name, (sheet.nrows, sheet.ncols), preview))
            workbook.unload_sheet(index)
        if not sheets:
            raise ValueError("The workbook contains no worksheets.")
        return tuple(sheets)
    finally:
        workbook.release_resources()


def _xls_records(sheet: "Sheet", region: TableRange, header: HeaderMode) -> Iterator[tuple[int, tuple[object, ...]]]:
    """Yield selected legacy cells, preserving gaps and rejecting date/error cells."""
    import xlrd
    last_row = region.row_end or sheet.nrows
    last_column = region.column_end or sheet.ncols
    for row in range(region.row_start - 1, last_row):
        cells: list[object] = [None] * (region.column_start - 1)
        skip_header = row == region.row_start - 1 and header == HeaderMode.FIRST
        for column in range(region.column_start - 1, last_column):
            kind = sheet.cell_type(row, column)
            value = sheet.cell_value(row, column)
            if kind in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                cells.append(None)
            elif skip_header or kind in (xlrd.XL_CELL_NUMBER, xlrd.XL_CELL_TEXT):
                cells.append(value)
            elif kind == xlrd.XL_CELL_BOOLEAN:
                cells.append(bool(value))
            else:
                reason = ("date/time cells are not supported as numeric matrix values"
                          if kind == xlrd.XL_CELL_DATE else
                          f"Excel error {xlrd.error_text_from_code.get(int(value), str(value))}"
                          if kind == xlrd.XL_CELL_ERROR else f"unsupported cell type {kind}")
                raise ValueError(f"Worksheet {sheet.name!r}, row {row + 1}, column {column + 1}: {reason}. Adjust the input range to exclude this cell.")
        yield row + 1, tuple(cells)


def _read_xls(path: Path, key: str | None, region: TableRange,
              header: HeaderMode) -> tuple[TableData, tuple[str, ...], str]:
    """Read one legacy numeric range; formulas use their saved result values."""
    region.validate()
    workbook = _open_xls(path)
    try:
        names = tuple(workbook.sheet_names())
        if not names or (key is not None and key not in names):
            raise ValueError(f"Worksheet {key!r} is unavailable. Worksheets: {', '.join(names)}")
        key = key if key is not None else names[0]
        sheet = workbook.sheet_by_name(key)
        if not sheet.nrows or not sheet.ncols:
            raise ValueError(f"Worksheet {key!r}: the selected range contains no numeric records.")
        if (region.row_start > sheet.nrows or region.column_start > sheet.ncols
                or (region.row_end is not None and region.row_end > sheet.nrows)
                or (region.column_end is not None and region.column_end > sheet.ncols)):
            raise ValueError(f"Worksheet {key!r}: selected range exceeds {sheet.nrows} rows × {sheet.ncols} columns.")
        table = numeric_table(_xls_records(sheet, region, header), region, header, f"Worksheet {key!r}")
        return table, names, key
    finally:
        workbook.release_resources()


def inspect_excel(path: Path) -> tuple[ExcelSheet, ...]:
    """List worksheet dimensions and the top-left preview without loading arrays.

    Args:
        path: XLSX, XLSM or legacy XLS workbook.

    Returns:
        Named worksheet rectangles, excluding chart-only sheets.

    Raises:
        ValueError: No worksheets; ImportError: Excel reader unavailable.
        OSError: Workbook cannot be opened. XML/ZIP errors propagate to the UI.
    """
    if path.suffix.lower() == ".xls":
        return _inspect_xls(path)
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheets: list[ExcelSheet] = []
        for source in workbook.worksheets:
            sheet = cast("ReadOnlyWorksheet", source)
            rows, columns = _dimensions(sheet)
            preview = tuple(tuple("" if value is None else str(value) for value in row)
                            for row in sheet.iter_rows(max_row=min(rows, PREVIEW_ROWS),
                                max_col=min(columns, PREVIEW_COLUMNS), values_only=True))
            sheets.append(ExcelSheet(sheet.title, (rows, columns), preview))
        if not sheets:
            raise ValueError("The workbook contains no worksheets.")
        return tuple(sheets)
    finally:
        workbook.close()


def read_excel(path: Path, key: str | None = None, region: TableRange = TableRange(),
               header: HeaderMode = HeaderMode.AUTO) -> tuple[TableData, tuple[str, ...], str]:
    """Read one sheet's numeric rectangle, validating cached formula values.

    Args:
        path: XLSX/XLSM/XLS workbook; source file is never written.
        key: Worksheet name, defaulting to the first worksheet.
        region: Original one-based inclusive row/column bounds.
        header: Treatment of the first selected record.

    Returns:
        Numeric table, worksheet names and selected name. Original coordinates
        appear in cell errors; matrix indices start at zero after import.

    Raises:
        ValueError: Bad range, missing sheet/cache, nonnumeric cells or precision
            loss. ImportError/OSError: Missing library or inaccessible workbook.
    """
    if path.suffix.lower() == ".xls":
        return _read_xls(path, key, region, header)
    from openpyxl import load_workbook
    region.validate()
    formulas = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        names = tuple(sheet.title for sheet in formulas.worksheets)
        if not names or (key is not None and key not in names):
            raise ValueError(f"Worksheet {key!r} is unavailable. Worksheets: {', '.join(names)}")
        key = key if key is not None else names[0]
        source = cast("ReadOnlyWorksheet", formulas[key])
        rows, columns = _dimensions(source)
        last_row, last_column = region.row_end or rows, region.column_end or columns
        if region.row_start > rows or region.column_start > columns or last_row > rows or last_column > columns:
            raise ValueError(f"Worksheet {key!r}: selected range exceeds {rows} rows × {columns} columns.")
        cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            values_sheet = cast("ReadOnlyWorksheet", cached[key])

            def records() -> Iterator[tuple[int, tuple[object, ...]]]:
                original = source.iter_rows(min_row=region.row_start, max_row=last_row, max_col=last_column)
                values = values_sheet.iter_rows(min_row=region.row_start, max_row=last_row, max_col=last_column, values_only=True)
                for number, (cells, row) in enumerate(zip(original, values, strict=True), region.row_start):
                    if not (number == region.row_start and header == HeaderMode.FIRST):
                        for column in range(region.column_start - 1, len(cells)):
                            if cells[column].data_type == "f" and row[column] is None:
                                raise ValueError(f"Worksheet {key!r}, row {number}, column {column + 1}: formula has no cached result. Recalculate and save in Excel/LibreOffice, or exclude this cell.")
                    yield number, row

            table = numeric_table(records(), region, header, f"Worksheet {key!r}")
            return table, names, key
        finally:
            cached.close()
    finally:
        formulas.close()


def excel_sheet_names(names: Sequence[str]) -> tuple[str, ...]:
    """Return unique legal 31-character sheet names, preserving useful aliases."""
    result: list[str] = []
    used: set[str] = set()
    for label in names:
        base = re.sub(r"[\\/*?:\[\]\x00-\x1f]", "_", label).strip(" '") or "matrix"
        name, number = base[:31], 2
        while name.casefold() in used:
            suffix = f"_{number}"
            name = f"{base[:31-len(suffix)]}{suffix}"
            number += 1
        result.append(name)
        used.add(name.casefold())
    return tuple(result)


def serialize_excel(matrices: Sequence[tuple[str, NDArray[np.generic]]]) -> bytes:
    """Save independent numeric matrices as worksheets in one new XLSX file.

    Args:
        matrices: Names and 1D/2D arrays. Vectors become columns; complex arrays
            become separate real/imaginary sheets. No formulas are generated.

    Returns:
        Workbook bytes. NaN/Inf and integers longer than 15 decimal digits use
        numeric text to preserve values on reimport; floats use Excel doubles.

    Raises:
        ValueError: No matrices, unsupported dtype/rank/precision or Excel limits.
        ImportError: openpyxl unavailable. Source arrays are never modified.
    """
    from openpyxl import Workbook
    sheets: list[tuple[str, NDArray[np.generic]]] = []
    for name, values in matrices:
        if values.ndim not in (1, 2) or values.dtype.kind not in "buifc" or not values.size:
            raise ValueError("Excel export requires a nonempty numeric 1D/2D array. Select a scalar image channel or use NPY.")
        if (values.dtype.kind == "f" and values.dtype.itemsize > 8
                or values.dtype.kind == "c" and values.dtype.itemsize > 16):
            raise ValueError("Excel cannot preserve extended floating-point precision. Use NPY.")
        shape = (values.shape[0], 1) if values.ndim == 1 else values.shape
        if shape[0] > EXCEL_MAX_ROWS or shape[1] > EXCEL_MAX_COLUMNS:
            raise ValueError(f"Shape {shape} exceeds Excel's {EXCEL_MAX_ROWS} row / {EXCEL_MAX_COLUMNS} column limits. Crop or use NPY/MAT.")
        if np.iscomplexobj(values):
            sheets.extend(((f"{name[:26]}_real", values.real), (f"{name[:26]}_imag", values.imag)))
        else:
            sheets.append((name, values))
    if not sheets:
        raise ValueError("No matrices to export.")
    workbook = Workbook(write_only=True)
    try:
        for title, (_, values) in zip(excel_sheet_names([name for name, _ in sheets]), sheets, strict=True):
            sheet = workbook.create_sheet(title)
            table = values[:, None] if values.ndim == 1 else values
            for row in table:
                cells: list[int | float | bool | str] = []
                for value in row:
                    if values.dtype.kind in "bui":
                        integer = int(value)
                        cells.append(str(integer) if abs(integer) >= 10 ** 15 else integer)
                    else:
                        number = float(value)
                        cells.append(number if np.isfinite(number) else str(number))
                sheet.append(cells)
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
    finally:
        workbook.close()
